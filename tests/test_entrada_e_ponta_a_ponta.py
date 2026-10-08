from openpyxl import Workbook, load_workbook

from matching import avaliacao, casamento, entrada, gerar_dados, saida


def _planilha(caminho, cabecalho, linhas):
    livro = Workbook()
    livro.active.append(cabecalho)
    for l in linhas:
        livro.active.append(l)
    livro.save(caminho)


def test_os_dois_formatos_viram_a_mesma_linha(tmp_path):
    _planilha(tmp_path / "obra.xlsx", ["ID_LINHA", "TITULO", "AUTORES", "INTERPRETES", "EDITORAS"],
              [["L1", "Saudade", "Ana Lima; Duarte, Pedro", "Olga Macedo", "Selo Maré"]])
    _planilha(tmp_path / "participante.xlsx", ["ID_LINHA", "TITULO", "PARTICIPANTE", "FUNCAO"],
              [["L1", "Saudade", "Ana Lima", "AUTOR"], ["L1", "Saudade", "Duarte, Pedro", "AUTOR"],
               ["L1", "Saudade", "Olga Macedo", "INTERPRETE"], ["L1", "Saudade", "Selo Maré", "EDITORA"]])
    assert entrada.ler_planilha(tmp_path / "obra.xlsx") == entrada.ler_planilha(tmp_path / "participante.xlsx")


def test_staging_nao_aceita_a_mesma_linha_duas_vezes(tmp_path):
    staging = entrada.Staging(tmp_path / "staging.sqlite")
    linhas = [entrada.Linha("L1", "Saudade", ["Ana Lima"]), entrada.Linha("L2", "Aurora", ["Olga Macedo"])]
    assert staging.receber("lote.xlsx", linhas) == (2, 0)
    assert staging.receber("lote_reenviado.xlsx", linhas) == (0, 2)
    assert staging.receber("lote_corrigido.xlsx", [entrada.Linha("L2", "Aurora", ["Olga Macedo", "Ana Lima"])]) == (1, 0)
    assert len(staging.todas()) == 3
    staging.fechar()


def test_ponta_a_ponta_contra_o_gabarito(tmp_path):
    gerar_dados.gerar(tmp_path)
    staging = entrada.Staging(tmp_path / "staging.sqlite")
    for lote in ("lote_1.xlsx", "lote_2.xlsx"):
        staging.receber(lote, entrada.ler_planilha(tmp_path / lote))
    linhas = staging.todas()
    staging.fechar()
    catalogo = casamento.ler_catalogo(tmp_path / "catalogo.csv")
    resultados = [casamento.casar(l, catalogo) for l in linhas]
    a = avaliacao.avaliar(resultados, tmp_path / "gabarito.csv")
    assert a.novas_casadas == 0                 # obra nova nunca casa sozinha com outra
    assert a.precisao >= 0.99
    assert a.existentes_pendentes == 0          # obra que existe no catálogo, no mínimo, vai para revisão

    saida.gravar(tmp_path / "resultado.xlsx", resultados, catalogo)
    livro = load_workbook(tmp_path / "resultado.xlsx", read_only=True)
    assert livro.sheetnames == ["Casadas", "Duvidas", "Pendentes", "Resumo"]
    resumo = {s: n for s, n in livro["Resumo"].iter_rows(min_row=2, values_only=True)}
    assert resumo["TOTAL"] == len(resultados) == resumo["CASADA"] + resumo["DUVIDA"] + resumo["PENDENTE"]
