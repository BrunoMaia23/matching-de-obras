"""Planilha de resultado para quem revisa: casadas, dúvidas com os três candidatos, pendentes e resumo."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .casamento import Catalogo, Resultado
from .decisao import CASADA, DUVIDA, PENDENTE


def gravar(caminho: Path | str, resultados: list[Resultado], catalogo: Catalogo) -> None:
    livro = Workbook()
    casadas = livro.active
    casadas.title = "Casadas"
    _aba(casadas, ["ID_LINHA", "TITULO_ENVIADO", "CODIGO", "TITULO_CATALOGO", "PONTOS", "PONTOS_TITULO",
                   "PONTOS_AUTORES"],
         [[r.linha.id_linha, r.linha.titulo, r.decisao.codigo, catalogo.obras[r.decisao.codigo]["titulo"],
           *_pontos(r.decisao.ranking[0][1])] for r in resultados if r.decisao.status == CASADA])

    duvidas = []
    for r in resultados:
        if r.decisao.status != DUVIDA:
            continue
        linha = [r.linha.id_linha, r.linha.titulo, "; ".join(r.linha.autores)]
        for codigo, p in r.decisao.ranking:
            obra = catalogo.obras[codigo]
            linha += [codigo, f"{obra['titulo']} ({'; '.join(obra['autores'])})", p.total]
        duvidas.append(linha)
    cabecalho = ["ID_LINHA", "TITULO_ENVIADO", "AUTORES_ENVIADOS"]
    for k in (1, 2, 3):
        cabecalho += [f"CANDIDATO_{k}", f"OBRA_{k}", f"PONTOS_{k}"]
    _aba(livro.create_sheet("Duvidas"), cabecalho, duvidas)

    _aba(livro.create_sheet("Pendentes"), ["ID_LINHA", "TITULO_ENVIADO", "AUTORES_ENVIADOS", "MELHOR_PONTUACAO"],
         [[r.linha.id_linha, r.linha.titulo, "; ".join(r.linha.autores),
           r.decisao.ranking[0][1].total if r.decisao.ranking else None]
          for r in resultados if r.decisao.status == PENDENTE])

    contagem = Counter(r.decisao.status for r in resultados)
    _aba(livro.create_sheet("Resumo"), ["SITUACAO", "LINHAS"],
         [[s, contagem.get(s, 0)] for s in (CASADA, DUVIDA, PENDENTE)] + [["TOTAL", len(resultados)]])
    livro.save(caminho)


def _pontos(p) -> list:
    return [p.total, p.titulo, p.autores]


def _aba(aba, cabecalho: list[str], linhas: list[list]) -> None:
    aba.append(cabecalho)
    for c in aba[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        aba.append(linha)
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    for i, nome in enumerate(cabecalho, start=1):
        largura = max([len(nome)] + [len(str(l[i - 1] or "")) for l in linhas[:200] if len(l) >= i])
        aba.column_dimensions[aba.cell(row=1, column=i).column_letter].width = min(60, largura + 2)
