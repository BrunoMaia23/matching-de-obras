"""Leitura das planilhas do parceiro e controle do que já foi recebido.

O parceiro manda em dois formatos: uma linha por obra (participantes em listas separadas por ";") ou
uma linha por participante. O segundo é pivotado para o primeiro. Cada linha ganha um hash do
conteúdo, e a staging só aceita hash novo: a mesma planilha reenviada não entra duas vezes.
"""
from __future__ import annotations

import hashlib
import json
import sqlite3
from dataclasses import asdict, dataclass, field
from pathlib import Path

from openpyxl import load_workbook

FORMATO_OBRA = {"ID_LINHA", "TITULO", "AUTORES", "INTERPRETES", "EDITORAS"}
FORMATO_PARTICIPANTE = {"ID_LINHA", "TITULO", "PARTICIPANTE", "FUNCAO"}
FUNCOES = {"AUTOR": "autores", "INTERPRETE": "interpretes", "EDITORA": "editoras"}


@dataclass
class Linha:
    id_linha: str
    titulo: str
    autores: list[str] = field(default_factory=list)
    interpretes: list[str] = field(default_factory=list)
    editoras: list[str] = field(default_factory=list)

    def hash(self) -> str:
        return hashlib.sha256(json.dumps(asdict(self), ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def ler_planilha(caminho: Path | str) -> list[Linha]:
    planilha = load_workbook(caminho, read_only=True).active
    linhas = planilha.iter_rows(values_only=True)
    cabecalho = [str(c or "").strip().upper() for c in next(linhas)]
    registros = [dict(zip(cabecalho, valores)) for valores in linhas if any(v is not None for v in valores)]
    if FORMATO_OBRA <= set(cabecalho):
        return [Linha(str(r["ID_LINHA"]), str(r["TITULO"] or ""), _separar(r["AUTORES"]),
                      _separar(r["INTERPRETES"]), _separar(r["EDITORAS"])) for r in registros]
    if FORMATO_PARTICIPANTE <= set(cabecalho):
        return pivotar(registros)
    raise ValueError(f"cabeçalho não reconhecido em {Path(caminho).name}: {cabecalho}")


def pivotar(registros: list[dict]) -> list[Linha]:
    """Uma linha por participante -> uma linha por obra, na ordem em que as obras aparecem."""
    obras: dict[str, Linha] = {}
    for r in registros:
        linha = obras.setdefault(str(r["ID_LINHA"]), Linha(str(r["ID_LINHA"]), str(r["TITULO"] or "")))
        funcao = FUNCOES.get(str(r["FUNCAO"] or "").strip().upper())
        if funcao is None:
            raise ValueError(f"função desconhecida na linha {r['ID_LINHA']}: {r['FUNCAO']!r}")
        if r["PARTICIPANTE"]:
            getattr(linha, funcao).append(str(r["PARTICIPANTE"]).strip())
    return list(obras.values())


def _separar(valor) -> list[str]:
    return [p.strip() for p in str(valor or "").split(";") if p.strip()]


class Staging:
    def __init__(self, caminho: Path | str):
        self.con = sqlite3.connect(caminho)
        self.con.execute("""CREATE TABLE IF NOT EXISTS entrada (
            hash TEXT PRIMARY KEY, arquivo TEXT NOT NULL, id_linha TEXT NOT NULL, linha TEXT NOT NULL,
            recebida_em TEXT NOT NULL DEFAULT (datetime('now')))""")

    def receber(self, arquivo: Path | str, linhas: list[Linha]) -> tuple[int, int]:
        """Grava as linhas novas. Devolve (novas, já vistas)."""
        novas = 0
        with self.con:
            for l in linhas:
                cursor = self.con.execute(
                    "INSERT OR IGNORE INTO entrada (hash, arquivo, id_linha, linha) VALUES (?, ?, ?, ?)",
                    [l.hash(), Path(arquivo).name, l.id_linha, json.dumps(asdict(l), ensure_ascii=False)])
                novas += cursor.rowcount
        return novas, len(linhas) - novas

    def todas(self) -> list[Linha]:
        return [Linha(**json.loads(texto)) for (texto,) in
                self.con.execute("SELECT linha FROM entrada ORDER BY recebida_em, arquivo, id_linha")]

    def fechar(self) -> None:
        self.con.close()
